"""
Fan Coil Selection Tool
-----------------------
Step 1: HAP report (xlsx)  ->  new Excel (spaces, areas in m2, loads) for manual inputs
Step 2: filled Excel + fancoil_database.xlsx (next to this program)  ->  fan coil selection

Loads/capacities in BTU/hr, airflow in CFM.
"""
import os
import re
import sys
import shutil
import subprocess
import traceback

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

FT2_TO_M2 = 0.09290304
M3_TO_FT3 = 35.3147
DB_NAME = "fancoil_database.xlsx"
TYPES = ["Concealed", "Ducted", "Cassette"]
SPEEDS = ["Low", "Medium", "High"]
# 1-based column of "Total Cooling" for each speed in the database sheets
# (+1 = sensible, +2 = heating, +3 = airflow CFM)
SPEED_COL = {"Low": 3, "Medium": 7, "High": 11}
DB_FIRST_ROW = 7
MAIN = "Fan Coil Selection"
FIRST = 5  # first data row in the selection sheet
HEADER_ROW = 4


# ----------------------------------------------------------------- helpers
def base_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def db_path():
    return os.path.join(base_dir(), DB_NAME)


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def tofloat(x):
    if num(x):
        return float(x)
    try:
        s = str(x).strip().replace(",", "")
        return float(s) if s else None
    except Exception:
        return None


def open_file(path):
    if sys.platform.startswith("win"):
        os.startfile(path)  # noqa
    elif sys.platform == "darwin":
        subprocess.call(["open", path])
    else:
        subprocess.call(["xdg-open", path])


# ------------------------------------------------------------ HAP report
RE_LOAD = re.compile(r'Component Loads For Space\s+"([^"]+)"\s+In Zone\s+"([^"]+)"')


def parse_hap(path):
    """Returns (spaces, warnings). Each space: name, zone, area_ft2, cool_sens, cool_lat, heat."""
    wb = load_workbook(path, read_only=True, data_only=True)
    areas, order, loads, warnings = {}, [], {}, []
    for ws in wb.worksheets:
        in_table = False
        cur = None
        for raw in ws.iter_rows(values_only=True):
            r = [c for c in raw if c is not None and str(c).strip() != ""]
            if not r:
                continue
            t = str(r[0]).strip()
            if t.startswith("Zone Name / Space Name"):
                in_table = True
                continue
            if in_table:
                if t.startswith("Totals"):
                    in_table = False
                    continue
                if len(r) >= 3 and num(r[1]) and num(r[2]):
                    if t not in areas:
                        order.append(t)
                    areas[t] = (float(r[1]), float(r[2]))
                    continue
            m = RE_LOAD.search(t)
            if m:
                cur = m.group(1)
                loads[cur] = {"zone": m.group(2)}
                continue
            if cur and t.startswith(">> Total Zone Loads"):
                vals = r[1:]
                if len(vals) >= 6:
                    sens, lat, heat = vals[1], vals[2], vals[4]
                else:
                    nums = [v for v in vals if num(v)]
                    if len(nums) != 4:
                        warnings.append(f"Could not read total loads of space {cur}")
                        cur = None
                        continue
                    sens, lat, heat = nums[0], nums[1], nums[2]
                loads[cur].update(
                    sens=float(sens) if num(sens) else 0.0,
                    lat=float(lat) if num(lat) else 0.0,
                    heat=float(heat) if num(heat) else 0.0,
                )
                cur = None
    spaces = []
    for name in order:
        mult, area = areas[name]
        if name not in loads or "sens" not in loads[name]:
            warnings.append(f"No load table found for space {name} (skipped)")
            continue
        if abs(mult - 1) > 1e-9:
            warnings.append(f"Space {name} has multiplier {mult:g}; loads are used as reported per space")
        z = loads[name]["zone"]
        spaces.append(
            dict(
                name=name,
                zone=("Zone " + z) if z.isdigit() else z,
                area_ft2=area,
                sens=loads[name]["sens"],
                lat=loads[name]["lat"],
                heat=loads[name]["heat"],
            )
        )
    for name in loads:
        if name not in areas:
            warnings.append(f"Space {name} has loads but no floor area in the first table (skipped)")
    return spaces, warnings


# -------------------------------------------------------------- database
def read_database(path):
    wb = load_workbook(path, data_only=True)
    db = {}
    for t in TYPES:
        models = []
        if t in wb.sheetnames:
            for r in wb[t].iter_rows(min_row=DB_FIRST_ROW, values_only=True):
                if len(r) < 2 or r[0] is None or r[1] is None:
                    continue
                brand, model = str(r[0]).strip(), str(r[1]).strip()
                if not brand or not model:
                    continue
                d = {"brand": brand, "model": model, "speeds": {}}
                for sp, c in SPEED_COL.items():
                    vals = []
                    for k in range(4):
                        idx = c - 1 + k
                        v = r[idx] if idx < len(r) else None
                        vals.append(float(v) if num(v) else None)
                    d["speeds"][sp] = tuple(vals)  # total, sensible, heating, cfm
                models.append(d)
        db[t] = models
    return db


def brand_list(db):
    seen, out = set(), []
    for t in TYPES:
        for m in db.get(t, []):
            k = m["brand"].casefold()
            if k not in seen:
                seen.add(k)
                out.append(m["brand"])
    return out


# -------------------------------------------------------------- altitude
ALT_TABLE = [  # ft, total cooling, sensible cooling, heating
    (0, 1.00, 1.00, 1.00),
    (1000, 0.99, 0.96, 0.96),
    (2000, 0.98, 0.93, 0.93),
    (3000, 0.97, 0.90, 0.90),
    (4000, 0.96, 0.86, 0.86),
    (5000, 0.94, 0.83, 0.83),
    (6000, 0.93, 0.80, 0.80),
    (7000, 0.92, 0.77, 0.77),
]


def read_altitude_table(ws):
    rows = []
    for r in ws.iter_rows(min_row=4, max_col=4, values_only=True):
        if num(r[0]) and all(num(v) for v in r[1:4]):
            f = [float(v) for v in r[1:4]]
            f = [v / 100 if v > 1.5 else v for v in f]
            rows.append((float(r[0]), *f))
    rows.sort()
    return rows


def interp_factors(rows, alt):
    """returns (total cooling, sensible cooling, heating) factors"""
    if not rows:
        raise ValueError("The Altitude Correction table is empty")
    if alt < rows[0][0] or alt > rows[-1][0]:
        raise ValueError(
            f"Altitude {alt:g} ft is outside the correction table ({rows[0][0]:g} to {rows[-1][0]:g} ft)"
        )
    if len(rows) == 1:
        return rows[0][1:]
    for i in range(len(rows) - 1):
        a0, a1 = rows[i][0], rows[i + 1][0]
        if a0 <= alt <= a1:
            t = (alt - a0) / (a1 - a0)
            return tuple(rows[i][k] + t * (rows[i + 1][k] - rows[i][k]) for k in (1, 2, 3))
    return rows[-1][1:]


# ------------------------------------------------- selection workbook
COLS = [
    "Zone", "Space", "Floor Area (ft²)", "Floor Area (m²)", "Height (m)", "Volume (m³)", "ACH",
    "Required Airflow (CFM)", "Cooling Sensible (BTU/hr)", "Cooling Latent (BTU/hr)",
    "Cooling Total (BTU/hr)", "Heating (BTU/hr)", "No. of Fan Coils", "Fan Coil Type",
    "Sensible per Unit (BTU/hr)", "Total Cooling per Unit (BTU/hr)", "Heating per Unit (BTU/hr)",
    "Airflow per Unit (CFM)",
    "Model by Sensible", "Model by Total", "Model by Heating", "Model by ACH", "Final Selected Model",
    "Catalog Total Cooling (BTU/hr)", "Catalog Sensible (BTU/hr)", "Catalog Heating (BTU/hr)",
    "Catalog Airflow (CFM)",
    "Total Cooling Margin", "Sensible Margin", "Heating Margin", "Airflow Margin", "Status",
]
INPUT_COLS = {5, 7, 13, 14}
EXTRACTED_COLS = {3, 9, 10, 12}
RESULT_COLS = set(range(19, 33))
BLUE = "0000FF"


def _font(**k):
    return Font(name="Arial", size=10, **k)


def _interp_formula(col):
    A = "'Altitude Correction'!$A$4:$A$11"
    V = f"'Altitude Correction'!${col}$4:${col}$11"
    i = f"MIN(MATCH($E$3,{A},1),7)"
    return (
        '=IF($E$3="","Enter altitude",IF(OR($E$3<\'Altitude Correction\'!$A$4,'
        '$E$3>\'Altitude Correction\'!$A$11),"Out of table range",'
        f"INDEX({V},{i})+($E$3-INDEX({A},{i}))/(INDEX({A},{i}+1)-INDEX({A},{i}))"
        f"*(INDEX({V},{i}+1)-INDEX({V},{i}))))"
    )


def build_workbook(spaces, brands):
    wb = Workbook()
    w = wb.active
    w.title = MAIN
    thin = Side(style="thin", color="999999")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    grey = PatternFill("solid", fgColor="E7E6E6")
    n = len(spaces) + HEADER_ROW

    # header
    for j, h in enumerate(COLS, 1):
        c = w.cell(HEADER_ROW, j, h)
        c.font = _font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = bd
        c.fill = PatternFill(
            "solid", fgColor="7F7F7F" if j in RESULT_COLS else ("BF8F00" if j in INPUT_COLS else "1F4E78")
        )
    w.row_dimensions[HEADER_ROW].height = 45

    # rows
    for i, s in enumerate(spaces):
        r = FIRST + i
        vals = {
            1: s["zone"], 2: s["name"], 3: s["area_ft2"],
            4: f"=C{r}*Legend!$B$2",
            6: f'=IF(E{r}="","",D{r}*E{r})',
            8: f'=IF(OR(F{r}="",G{r}=""),"",F{r}*G{r}*Legend!$B$3/60)',
            9: s["sens"], 10: s["lat"], 11: f"=I{r}+J{r}", 12: s["heat"],
            15: f'=IF(N(M{r})>0,I{r}/M{r},"")',
            16: f'=IF(N(M{r})>0,K{r}/M{r},"")',
            17: f'=IF(N(M{r})>0,L{r}/M{r},"")',
            18: f'=IF(AND(N(M{r})>0,H{r}<>""),H{r}/M{r},"")',
            28: f'=IF(OR(X{r}="",N(P{r})=0),"",X{r}*$K$3/P{r}-1)',
            29: f'=IF(OR(Y{r}="",N(O{r})=0),"",Y{r}*$I$3/O{r}-1)',
            30: f'=IF(OR(Z{r}="",N(Q{r})=0),"",Z{r}*$L$3/Q{r}-1)',
            31: f'=IF(OR(AA{r}="",N(R{r})=0),"",AA{r}/R{r}-1)',
        }
        for j in range(1, len(COLS) + 1):
            c = w.cell(r, j, vals.get(j))
            c.border = bd
            c.font = _font(color=BLUE) if j in EXTRACTED_COLS else _font()
            if j in INPUT_COLS:
                c.fill = yellow
                c.font = _font(color=BLUE)
            if j in RESULT_COLS:
                c.fill = grey
            if j in (3, 4, 6):
                c.number_format = "#,##0.0"
            elif j in (8, 9, 10, 11, 12, 15, 16, 17, 18, 24, 25, 26, 27):
                c.number_format = "#,##0"
            elif j == 5:
                c.number_format = "0.00"
            elif j == 7:
                c.number_format = "0.0"
            elif 28 <= j <= 31:
                c.number_format = "0%"

    # settings block (rows 1-3)
    w["A1"] = "Fan Speed (all spaces):"
    w["A2"] = "Brand (all spaces):"
    w["A3"] = "Altitude above sea level (ft):"
    for a in ("A1", "A2", "A3"):
        w[a].font = _font(bold=True)
    w["E1"] = "Medium"
    for a in ("E1", "E2", "E3"):
        w[a].font = _font(bold=True, color=BLUE)
        w[a].fill = yellow
        w[a].border = bd
        w[a].alignment = Alignment(horizontal="center")
    w["E3"].number_format = "#,##0"
    w["F1"] = "← Low / Medium / High; catalog data of this speed is used for the selection"
    w["F2"] = "← brand name exactly as written in fancoil_database.xlsx"
    w["F3"] = "Capacity factor →"
    for a in ("F1", "F2", "F3"):
        w[a].font = _font(italic=True)
    for cell, col in (("I3", "C"), ("K3", "B"), ("L3", "D")):
        w[cell] = _interp_formula(col)
        w[cell].font = _font(bold=True)
        w[cell].border = bd
        w[cell].number_format = "0.0%"
        w[cell].alignment = Alignment(horizontal="center")

    # validations
    dv = DataValidation(type="list", formula1='"Concealed,Ducted,Cassette"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid type",
                        error="Choose Concealed, Ducted or Cassette")
    dv2 = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                         showErrorMessage=True, error="Enter a whole number >= 1")
    dv3 = DataValidation(type="list", formula1='"Low,Medium,High"', allow_blank=False,
                         showErrorMessage=True, errorTitle="Invalid speed", error="Choose Low, Medium or High")
    dv4 = DataValidation(type="list", formula1="=BrandList", allow_blank=True, showErrorMessage=False)
    dv5 = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                         showErrorMessage=True, error="Enter altitude in feet (>= 0)")
    for d_, rng in ((dv, f"N{FIRST}:N{n}"), (dv2, f"M{FIRST}:M{n}"), (dv3, "E1"), (dv4, "E2"), (dv5, "E3")):
        w.add_data_validation(d_)
        d_.add(rng)

    # conditional formats
    red = PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")
    w.conditional_formatting.add(
        f"AF{FIRST}:AF{n}", FormulaRule(formula=[f'AND($AF{FIRST}<>"",$AF{FIRST}<>"OK")'], fill=red))
    w.conditional_formatting.add(
        f"W{FIRST}:W{n}", FormulaRule(formula=[f'$W{FIRST}="NOT FOUND"'], fill=red))

    widths = [9, 11, 11, 11, 9, 11, 7, 14, 14, 14, 14, 13, 10, 14, 14, 15, 14, 13,
              15, 15, 15, 15, 18, 14, 14, 14, 14, 12, 12, 12, 12, 48]
    for j, wd in enumerate(widths, 1):
        w.column_dimensions[get_column_letter(j)].width = wd
    w.freeze_panes = f"C{FIRST}"

    # Lists
    ls = wb.create_sheet("Lists")
    ls["A1"] = "Brand list (filled by the program from fancoil_database.xlsx; typing in E2 is also allowed)"
    ls["A1"].font = _font(bold=True)
    ls.column_dimensions["A"].width = 40
    for i, b in enumerate(brands[:100]):
        ls.cell(2 + i, 1, b).font = _font(color=BLUE)
    dn = DefinedName("BrandList", attr_text="OFFSET(Lists!$A$2,0,0,MAX(1,COUNTA(Lists!$A$2:$A$101)),1)")
    try:
        wb.defined_names["BrandList"] = dn
    except TypeError:
        wb.defined_names.append(dn)

    # Altitude Correction
    ac = wb.create_sheet("Altitude Correction")
    ac["A1"] = "Altitude correction of equipment capacity (% of catalog capacity available at altitude)"
    ac["A1"].font = _font(bold=True)
    ac["A2"] = "Editable reference table. Intermediate altitudes are linearly interpolated."
    ac["A2"].font = _font(italic=True)
    for j, h in enumerate(["Altitude above sea level (ft)", "Total Cooling Capacity",
                           "Sensible Cooling Capacity", "Total Heating Capacity"], 1):
        c = ac.cell(3, j, h)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.border = bd
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(ALT_TABLE):
        for j, v in enumerate(row, 1):
            c = ac.cell(4 + i, j, v)
            c.font = _font(color=BLUE)
            c.border = bd
            c.number_format = "#,##0" if j == 1 else "0%"
    for col, wd in zip("ABCD", [22, 18, 18, 18]):
        ac.column_dimensions[col].width = wd
    ac.row_dimensions[3].height = 32
    ac["A13"] = "Source: table provided by the user; 0 ft = 100% added by the user"
    ac["A13"].font = _font(italic=True)

    # Legend
    lg = wb.create_sheet("Legend")
    lg["A1"], lg["B1"] = "Constant", "Value"
    lg["A2"], lg["B2"] = "ft² → m²", FT2_TO_M2
    lg["A3"], lg["B3"] = "m³ → ft³", M3_TO_FT3
    notes = [
        "Yellow cells = manual input (Fan speed, Brand, Altitude, Height, ACH, No. of Fan Coils, Type)",
        "Blue font = values extracted from the HAP report / catalog; black = formulas",
        "Grey columns = results written by the program (Step 2)",
        "Per-unit loads = space load ÷ No. of Fan Coils; Required Airflow = Volume × ACH × 35.3147 ÷ 60 (CFM)",
        "HAP loads already include the HAP safety factor; the altitude factors derate the CATALOG capacity",
        "Altitude (E3) is the project elevation in ft, not the room height",
        "Selection: smallest model (sorted by catalog total cooling at the chosen speed) where "
        "catalog capacity × altitude factor ≥ per-unit load, and catalog airflow ≥ per-unit required airflow",
        "Margin = (capacity after altitude factor ÷ per-unit load) − 1; airflow margin has no altitude factor",
    ]
    lg["A5"] = "Notes"
    lg["A5"].font = _font(bold=True)
    for i, t in enumerate(notes):
        lg.cell(6 + i, 1, t).font = _font()
    for a in ("A1", "B1"):
        lg[a].font = _font(bold=True)
    for a in ("A2", "A3"):
        lg[a].font = _font()
    for a in ("B2", "B3"):
        lg[a].font = _font(color=BLUE)
    lg.column_dimensions["A"].width = 24

    wb.calculation.fullCalcOnLoad = True
    return wb


# ------------------------------------------------------------- selection
def _first_ok(models, ok):
    for m in models:
        if ok(m):
            return m
    return None


def select_fancoils(in_path, out_path, database_path):
    """Reads the filled workbook, selects fan coils from the database, saves to out_path.
    Returns (summary_lines)."""
    wb = load_workbook(in_path)
    if MAIN not in wb.sheetnames or "Altitude Correction" not in wb.sheetnames:
        raise ValueError("This file is not a workbook created by this program (sheet not found).")
    w = wb[MAIN]
    if str(w.cell(HEADER_ROW, 2).value).strip() != "Space":
        raise ValueError("Unexpected layout: header row not found at row 4.")

    speed = str(w["E1"].value or "").strip().capitalize()
    brand = str(w["E2"].value or "").strip()
    alt = tofloat(w["E3"].value)
    if speed not in SPEEDS:
        raise ValueError("Choose the fan speed in cell E1 (Low / Medium / High).")
    if not brand:
        raise ValueError("Choose the brand in cell E2.")
    if alt is None:
        raise ValueError("Enter the altitude above sea level (ft) in cell E3.")
    f_tot, f_sens, f_heat = interp_factors(read_altitude_table(wb["Altitude Correction"]), alt)

    db = read_database(database_path)
    brands = brand_list(db)
    ls = wb["Lists"] if "Lists" in wb.sheetnames else wb.create_sheet("Lists")
    for r in range(2, 102):
        ls.cell(r, 1).value = None
    for i, b in enumerate(brands[:100]):
        ls.cell(2 + i, 1, b)

    by_type = {}
    for t in TYPES:
        ms = [m for m in db.get(t, []) if m["brand"].casefold() == brand.casefold()]
        ms.sort(key=lambda m: ((m["speeds"][speed][0] or 0), (m["speeds"][speed][1] or 0)))
        by_type[t] = ms

    summary = [f"Speed: {speed} | Brand: {brand} | Altitude: {alt:g} ft",
               f"Altitude factors -> total cooling {f_tot:.4f}, sensible {f_sens:.4f}, heating {f_heat:.4f}"]
    ok_count = warn_count = 0
    r = FIRST
    while w.cell(r, 2).value not in (None, ""):
        for c in list(range(19, 28)) + [32]:  # clear previous results (margin columns keep their formulas)
            w.cell(r, c).value = None
        name = w.cell(r, 2).value
        area_ft2 = tofloat(w.cell(r, 3).value)
        height = tofloat(w.cell(r, 5).value)
        ach = tofloat(w.cell(r, 7).value)
        sens = tofloat(w.cell(r, 9).value) or 0.0
        lat = tofloat(w.cell(r, 10).value) or 0.0
        heat = tofloat(w.cell(r, 12).value) or 0.0
        count = tofloat(w.cell(r, 13).value)
        typ = str(w.cell(r, 14).value or "").strip().capitalize()

        def status(msg, ok=False):
            nonlocal ok_count, warn_count
            w.cell(r, 32).value = msg
            if ok:
                ok_count += 1
            else:
                warn_count += 1

        if not count or count < 1 or typ not in TYPES:
            status("Enter No. of Fan Coils and Type")
            r += 1
            continue
        models = by_type[typ]
        if not models:
            status(f"No data for brand '{brand}' in {typ} database")
            r += 1
            continue
        count = int(round(count))
        tot_u = (sens + lat) / count
        sens_u = sens / count
        heat_u = heat / count
        air_u = None
        if area_ft2 is not None and height and ach:
            air_u = area_ft2 * FT2_TO_M2 * height * ach * M3_TO_FT3 / 60 / count

        def cap(m, k):
            return m["speeds"][speed][k] or 0.0

        by_s = _first_ok(models, lambda m: cap(m, 1) * f_sens >= sens_u - 1e-9)
        by_t = _first_ok(models, lambda m: cap(m, 0) * f_tot >= tot_u - 1e-9)
        by_h = _first_ok(models, lambda m: cap(m, 2) * f_heat >= heat_u - 1e-9)
        by_a = _first_ok(models, lambda m: cap(m, 3) >= air_u - 1e-9) if air_u else None
        final = _first_ok(
            models,
            lambda m: cap(m, 1) * f_sens >= sens_u - 1e-9
            and cap(m, 0) * f_tot >= tot_u - 1e-9
            and cap(m, 2) * f_heat >= heat_u - 1e-9
            and (air_u is None or cap(m, 3) >= air_u - 1e-9),
        )
        nm = lambda m: m["model"] if m else "None"
        w.cell(r, 19).value = nm(by_s)
        w.cell(r, 20).value = nm(by_t)
        w.cell(r, 21).value = nm(by_h)
        w.cell(r, 22).value = nm(by_a) if air_u else "— (no ACH)"
        if final:
            w.cell(r, 23).value = final["model"]
            w.cell(r, 24).value = cap(final, 0)
            w.cell(r, 25).value = cap(final, 1)
            w.cell(r, 26).value = cap(final, 2)
            w.cell(r, 27).value = cap(final, 3)
            msg = "OK" if air_u else "OK (airflow not checked: height/ACH missing)"
            status(msg, ok=True)
        else:
            w.cell(r, 23).value = "NOT FOUND"
            status("No single model satisfies all conditions - increase No. of Fan Coils or change type/brand")
        r += 1

    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    summary.append(f"Spaces processed: {ok_count + warn_count} | selected: {ok_count} | need attention: {warn_count}")
    return summary


# ------------------------------------------------------------------- GUI
def main():
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("Fan Coil Selection Tool")
    root.geometry("760x560")
    hap_var = tk.StringVar()
    sel_var = tk.StringVar()

    log = scrolledtext.ScrolledText(root, height=16, state="disabled", font=("Consolas", 9))

    def say(msg):
        log.configure(state="normal")
        log.insert("end", msg + "\n")
        log.see("end")
        log.configure(state="disabled")
        root.update_idletasks()

    def browse(var, title):
        p = filedialog.askopenfilename(title=title, filetypes=[("Excel", "*.xlsx *.xlsm")])
        if p:
            var.set(p)

    def need_db():
        p = db_path()
        if not os.path.exists(p):
            raise FileNotFoundError(f"Database file not found next to the program:\n{p}")
        return p

    def step1():
        try:
            src = hap_var.get().strip()
            if not src:
                raise ValueError("Select the HAP report first.")
            spaces, warns = parse_hap(src)
            if not spaces:
                raise ValueError("No spaces were found in the report.")
            brands = []
            if os.path.exists(db_path()):
                brands = brand_list(read_database(db_path()))
            stem = os.path.splitext(os.path.basename(src))[0]
            out = filedialog.asksaveasfilename(
                title="Save the new Excel as", defaultextension=".xlsx",
                initialdir=os.path.dirname(src), initialfile=stem + "_fancoil.xlsx",
                filetypes=[("Excel", "*.xlsx")])
            if not out:
                return
            build_workbook(spaces, brands).save(out)
            say(f"Created: {out}")
            say(f"{len(spaces)} spaces extracted; brands in database: {', '.join(brands) or '-'}")
            for wmsg in warns:
                say("WARNING: " + wmsg)
            sel_var.set(out)
            say("Next: open the file, fill the yellow cells, save, then press 'Select fan coils'.")
            if messagebox.askyesno("Open", "Open the new Excel file now?"):
                open_file(out)
        except Exception as e:  # noqa
            say("ERROR: " + str(e))
            messagebox.showerror("Error", str(e))

    def step2():
        try:
            src = sel_var.get().strip()
            if not src:
                raise ValueError("Select the filled Excel file first.")
            dbp = need_db()
            stem = os.path.splitext(os.path.basename(src))[0]
            out = filedialog.asksaveasfilename(
                title="Save the result as", defaultextension=".xlsx",
                initialdir=os.path.dirname(src), initialfile=stem + "_selected.xlsx",
                filetypes=[("Excel", "*.xlsx")])
            if not out:
                return
            for line in select_fancoils(src, out, dbp):
                say(line)
            say(f"Saved: {out}")
            if messagebox.askyesno("Open", "Open the result now?"):
                open_file(out)
        except PermissionError:
            msg = "A file is locked. Close the Excel files (database / input / output) and try again."
            say("ERROR: " + msg)
            messagebox.showerror("File is open", msg)
        except Exception as e:  # noqa
            say("ERROR: " + str(e))
            say(traceback.format_exc())
            messagebox.showerror("Error", str(e))

    def open_db():
        try:
            open_file(need_db())
        except Exception as e:  # noqa
            messagebox.showerror("Error", str(e))

    pad = {"padx": 10, "pady": 5}
    f1 = tk.LabelFrame(root, text="Step 1 - HAP report -> new Excel")
    f1.pack(fill="x", **pad)
    tk.Entry(f1, textvariable=hap_var, width=70).grid(row=0, column=0, padx=6, pady=6)
    tk.Button(f1, text="Browse...", command=lambda: browse(hap_var, "Select HAP report")).grid(row=0, column=1)
    tk.Button(f1, text="Create Excel from HAP report", command=step1).grid(row=1, column=0, columnspan=2, pady=4)

    f2 = tk.LabelFrame(root, text="Step 2 - filled Excel + database -> fan coil selection")
    f2.pack(fill="x", **pad)
    tk.Entry(f2, textvariable=sel_var, width=70).grid(row=0, column=0, padx=6, pady=6)
    tk.Button(f2, text="Browse...", command=lambda: browse(sel_var, "Select the filled Excel")).grid(row=0, column=1)
    tk.Button(f2, text="Select fan coils", command=step2).grid(row=1, column=0, columnspan=2, pady=4)

    tk.Button(root, text="Database (open fancoil_database.xlsx)", command=open_db).pack(**pad)
    log.pack(fill="both", expand=True, **pad)
    say("Database file: " + db_path())
    root.mainloop()


if __name__ == "__main__":
    main()
